#!/usr/bin/env python3
"""Regenerate the Word fixture documents used to exercise the VBA macros.

Run from the repository root:

    python tests/fixtures/vba/make_fixtures.py

Outputs (kept next to this script):
    macro_smoke.docx       - literal $$...$$ / $...$ text for ConvertAllLatex
    renumber_targets.docx  - app-style numbered equations for Renumber/
                             InsertEquationReference tests
"""

from pathlib import Path

from docx import Document

import sys
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[2]))

from core.converter import LatexToWordConverter  # noqa: E402

OUT_DIR = HERE


def make_macro_smoke():
    """Document with RAW LaTeX text — macros must do the converting."""
    doc = Document()
    doc.add_heading("Macro Smoke Fixture", level=1)

    doc.add_paragraph(
        "Inline math $x^2 + y^2 = z^2$ and Greek $\\alpha + \\beta = \\gamma$."
    )
    doc.add_paragraph(
        "Currency guard: the price is $5.00 and the other is $10.00 today."
    )
    doc.add_paragraph(
        "Escaped dollars: the cost is \\$5 total."
    )

    doc.add_paragraph("Multiline display equation:")
    doc.add_paragraph("$$")
    doc.add_paragraph("E = mc^2")
    doc.add_paragraph("$$")

    doc.add_paragraph("Single-line display equation:")
    doc.add_paragraph("$$a^2 + b^2 = c^2$$")

    path = OUT_DIR / "macro_smoke.docx"
    doc.save(path)
    print(f"wrote {path}")


def make_renumber_targets():
    """App-generated numbered equation tables ((1), (2), (3))."""
    md = """# Renumber Targets

$$
E = mc^2
$$

$$
F = ma
$$

$$
p = mv
$$
"""
    conv = LatexToWordConverter()
    path = OUT_DIR / "renumber_targets.docx"
    stats = conv.convert_markdown_to_docx(md, str(path), True, "({n})")
    print(f"wrote {path} ({stats['equations_converted']} equations)")


if __name__ == "__main__":
    make_macro_smoke()
    make_renumber_targets()
